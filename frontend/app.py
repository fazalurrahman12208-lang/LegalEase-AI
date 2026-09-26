import os
import io
import requests
import streamlit as st
from dotenv import load_dotenv
from docx import Document
from fpdf import FPDF


# --------------------------------------------------
# Configuration
# --------------------------------------------------

load_dotenv()

BACKEND_URL = os.getenv(
    "BACKEND_URL",
    "http://127.0.0.1:8000"
)

GENERATE_URL = f"{BACKEND_URL}/generate"


# --------------------------------------------------
# Page Configuration
# --------------------------------------------------

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)


# --------------------------------------------------
# Custom Styling
# --------------------------------------------------

st.markdown(
    """
    <style>

    .main-title {
        text-align: center;
        font-size: 42px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        color: #777;
        font-size: 18px;
        margin-bottom: 30px;
    }

    .preview-box {
        background-color: #111827;
        color: white;
        padding: 25px;
        border-radius: 12px;
        min-height: 400px;
        max-height: 600px;
        overflow-y: auto;
        white-space: pre-wrap;
        font-family: Georgia, serif;
        line-height: 1.7;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# --------------------------------------------------
# Header
# --------------------------------------------------

st.markdown(
    '<div class="main-title">⚖️ LegalEase</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">AI-Powered Legal Document Generator</div>',
    unsafe_allow_html=True
)


# --------------------------------------------------
# Session State
# --------------------------------------------------

if "generated_document" not in st.session_state:
    st.session_state.generated_document = ""

if "document_type" not in st.session_state:
    st.session_state.document_type = ""


# --------------------------------------------------
# Input Section
# --------------------------------------------------

st.header("Create Your Legal Document")

col1, col2 = st.columns(2)

with col1:

    document_type = st.text_input(
        "Document Type",
        placeholder="Example: Freelance Work Contract",
        help="Enter the type of legal document you want to generate."
    )

    parties = st.text_area(
        "Parties Involved",
        placeholder=(
            "Example:\n"
            "Jane Doe (Service Provider), "
            "TechNova Inc. (Client)"
        ),
        height=120
    )


with col2:

    dates = st.text_input(
        "Effective Date",
        placeholder="Example: September 25, 2026"
    )

    terms = st.text_area(
        "Terms & Conditions",
        placeholder=(
            "Enter terms separated by semicolons.\n\n"
            "Example:\n"
            "Payment within 30 days; "
            "Confidentiality must be maintained; "
            "Either party may terminate with 15 days notice"
        ),
        height=120
    )


# --------------------------------------------------
# Generate Document
# --------------------------------------------------

st.divider()

generate_button = st.button(
    "Generate Document",
    type="primary",
    use_container_width=True
)


if generate_button:

    if not document_type.strip():
        st.error("Please enter a document type.")

    elif not parties.strip():
        st.error("Please enter the parties involved.")

    elif not terms.strip():
        st.error("Please enter the terms and conditions.")

    elif not dates.strip():
        st.error("Please enter the effective date.")

    else:

        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "effective_date": dates
        }

        try:

            with st.spinner(
                "Generating your legal document..."
            ):

                response = requests.post(
                    GENERATE_URL,
                    json=payload,
                    timeout=120
                )

            if response.status_code == 200:

                data = response.json()

                st.session_state.generated_document = (
                    data.get("document", "")
                )

                st.session_state.document_type = document_type

                st.success(
                    "Legal document generated successfully!"
                )

            else:

                try:
                    error_data = response.json()
                    error_message = error_data.get(
                        "detail",
                        "Unknown backend error."
                    )
                except Exception:
                    error_message = response.text

                st.error(
                    f"Backend error ({response.status_code}): "
                    f"{error_message}"
                )

        except requests.exceptions.ConnectionError:

            st.error(
                "Could not connect to the FastAPI backend. "
                "Make sure the backend is running on "
                f"{BACKEND_URL}."
            )

        except requests.exceptions.Timeout:

            st.error(
                "The request timed out. "
                "Please try again."
            )

        except Exception as exc:

            st.error(
                f"Unexpected error: {exc}"
            )


# --------------------------------------------------
# Generated Document
# --------------------------------------------------

if st.session_state.generated_document:

    st.divider()

    st.header("Generated Document")

    # --------------------------------------------------
    # Editable Document
    # --------------------------------------------------

    edited_document = st.text_area(
        "Edit Document",
        value=st.session_state.generated_document,
        height=600,
        help="You can edit the generated document before downloading."
    )

    st.session_state.generated_document = edited_document

    # --------------------------------------------------
    # Preview
    # --------------------------------------------------

    st.subheader("Document Preview")

    preview_text = (
        st.session_state.generated_document
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br>")
    )

    st.markdown(
        f"""
        <div class="preview-box">
            {preview_text}
        </div>
        """,
        unsafe_allow_html=True
    )

    # --------------------------------------------------
    # Download Functions
    # --------------------------------------------------

    document_text = st.session_state.generated_document
    safe_document_type = (
        st.session_state.document_type
        .strip()
        .replace(" ", "_")
        .lower()
    )

    # ==================================================
    # TXT
    # ==================================================

    txt_data = document_text.encode("utf-8")

    # ==================================================
    # DOCX
    # ==================================================

    doc = Document()

    title = doc.add_heading(
        st.session_state.document_type,
        level=0
    )

    doc.add_paragraph("Generated by LegalEase")

    doc.add_paragraph("")

    for paragraph in document_text.split("\n"):

        if paragraph.strip():

            doc.add_paragraph(
                paragraph.strip()
            )

    docx_buffer = io.BytesIO()

    doc.save(docx_buffer)

    docx_buffer.seek(0)

    # ==================================================
    # PDF
    # ==================================================

    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    pdf.set_font(
        "Times",
        "B",
        18
    )

    pdf.multi_cell(
        0,
        10,
        st.session_state.document_type,
        align="C"
    )

    pdf.ln(5)

    pdf.set_font(
        "Times",
        "",
        11
    )

    # FPDF's core fonts don't reliably support
    # all Unicode characters, so normalize common
    # typographic characters.

    pdf_text = (
        document_text
        .replace("–", "-")
        .replace("—", "-")
        .replace("“", '"')
        .replace("”", '"')
        .replace("‘", "'")
        .replace("’", "'")
        .replace("•", "-")
        .replace("…", "...")
    )

    for paragraph in pdf_text.split("\n"):

        paragraph = paragraph.strip()

        if paragraph:

            pdf.multi_cell(
                0,
                7,
                paragraph
            )

            pdf.ln(2)

    pdf.ln(8)

    pdf.set_font(
        "Times",
        "I",
        9
    )

    pdf.cell(
        0,
        10,
        "Generated by LegalEase",
        align="C"
    )

    pdf_bytes = bytes(pdf.output())

    # --------------------------------------------------
    # Download Buttons
    # --------------------------------------------------

    st.subheader("Download Document")

    download_col1, download_col2, download_col3 = st.columns(3)

    with download_col1:

        st.download_button(
            label="Download TXT",
            data=txt_data,
            file_name=f"{safe_document_type}.txt",
            mime="text/plain",
            use_container_width=True
        )

    with download_col2:

        st.download_button(
            label="Download DOCX",
            data=docx_buffer,
            file_name=f"{safe_document_type}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True
        )

    with download_col3:

        st.download_button(
            label="Download PDF",
            data=pdf_bytes,
            file_name=f"{safe_document_type}.pdf",
            mime="application/pdf",
            use_container_width=True
        )


# --------------------------------------------------
# Footer
# --------------------------------------------------

st.divider()

st.caption(
    "LegalEase • AI-Powered Legal Document Generator"
)